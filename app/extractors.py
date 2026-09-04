from __future__ import annotations

import csv
import bz2
import gzip
import html
import io
import json
import lzma
import mimetypes
import re
import shutil
import tarfile
import zipfile
from dataclasses import dataclass
from datetime import datetime
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from pathlib import Path
from typing import Callable, Iterable

from PIL import Image, ImageOps

from .errors import JobCancelled

try:
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:
    pass

from .media import (
    compress_video,
    extract_audio,
    extract_video_frames,
    human_size,
    media_summary,
    probe_media,
    transcribe_audio,
)
from .models import ProcessingOptions


AUDIO_EXTENSIONS = {
    ".mp3", ".wav", ".m4a", ".aac", ".flac", ".ogg", ".oga", ".opus", ".wma", ".aiff", ".aif", ".amr", ".ape",
}
VIDEO_EXTENSIONS = {
    ".mp4", ".mov", ".mkv", ".avi", ".webm", ".wmv", ".flv", ".m4v", ".mpeg", ".mpg", ".ts", ".mts", ".m2ts", ".3gp",
}
IMAGE_EXTENSIONS = {
    ".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff", ".ico", ".jfif", ".avif", ".heic", ".heif",
}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf", ".yaml", ".yml", ".toml", ".json", ".jsonl",
    ".xml", ".svg", ".py", ".js", ".jsx", ".ts", ".tsx", ".java", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs",
    ".php", ".rb", ".sh", ".ps1", ".bat", ".sql", ".css", ".scss", ".vue", ".svelte", ".tex", ".properties",
}
ARCHIVE_EXTENSIONS = {".zip", ".7z", ".tar", ".gz", ".tgz", ".bz2", ".xz"}

MAX_TEXT_CHARS = 3_000_000
MAX_ARCHIVE_BYTES = 4 * 1024 * 1024 * 1024


@dataclass
class VisualAsset:
    path: Path
    label: str
    source: str
    kind: str


class TextHTMLParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._hidden = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style", "noscript"}:
            self._hidden += 1
        if tag in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "h5", "tr"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style", "noscript"} and self._hidden:
            self._hidden -= 1

    def handle_data(self, data: str) -> None:
        if not self._hidden:
            self.parts.append(data)

    def text(self) -> str:
        value = html.unescape("".join(self.parts))
        return re.sub(r"\n\s*\n\s*\n+", "\n\n", value).strip()


class UniversalExtractor:
    def __init__(
        self,
        options: ProcessingOptions,
        job_dir: Path,
        status: Callable[[str, float | None], None],
    ) -> None:
        self.options = options
        self.job_dir = job_dir
        self.temp_dir = job_dir / "working"
        self.visual_dir = self.temp_dir / "visuals"
        self.archive_dir = self.temp_dir / "archives"
        self.audio_dir = self.temp_dir / "audio"
        self.compressed_work_dir = self.temp_dir / "compressed"
        self.compressed_keep_dir = job_dir / "compressed"
        self.temp_dir.mkdir(parents=True, exist_ok=True)
        self.visual_dir.mkdir(parents=True, exist_ok=True)
        self.status = status
        self.sections: list[tuple[str, str]] = []
        self.visuals: list[VisualAsset] = []
        self.warnings: list[str] = []
        self.failed_files: list[str] = []
        self.kept_compressed_files: list[Path] = []
        self.stats = {
            "files_seen": 0,
            "files_extracted": 0,
            "transcripts": 0,
            "visual_candidates": 0,
            "archives": 0,
            "compressed_videos": 0,
            "compressed_bytes_before": 0,
            "compressed_bytes_after": 0,
        }
        self._asset_counter = 0
        self._ocr_engine = None

    def extract(self, sources: Iterable[Path]) -> tuple[list[tuple[str, str]], list[VisualAsset], list[str], dict]:
        for source in sources:
            input_root = self.job_dir / "input"
            display_name = source.relative_to(input_root).as_posix() if source.is_relative_to(input_root) else source.name
            self.process_file(source, display_name, archive_depth=0)
        self.stats["visual_candidates"] = len(self.visuals)
        return self.sections, self.visuals, self.warnings, self.stats

    def process_file(self, path: Path, display_name: str, archive_depth: int) -> None:
        if not path.is_file():
            return
        self.stats["files_seen"] += 1
        self.status(f"正在读取 {display_name}")
        suffix = path.suffix.lower()

        if self.options.include_metadata:
            self.add_section(f"文件信息｜{display_name}", self.file_metadata(path))

        try:
            if suffix in ARCHIVE_EXTENSIONS or path.name.lower().endswith((".tar.gz", ".tar.bz2", ".tar.xz")):
                self.extract_archive(path, display_name, archive_depth)
            elif suffix in VIDEO_EXTENSIONS or suffix in AUDIO_EXTENSIONS:
                self.extract_media(path, display_name)
            elif suffix in IMAGE_EXTENSIONS:
                self.extract_image(path, display_name)
            elif suffix == ".pdf":
                self.extract_pdf(path, display_name)
            elif suffix == ".docx":
                self.extract_docx(path, display_name)
            elif suffix == ".pptx":
                self.extract_pptx(path, display_name)
            elif suffix in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
                self.extract_xlsx(path, display_name)
            elif suffix in {".csv", ".tsv"}:
                self.extract_delimited(path, display_name, "\t" if suffix == ".tsv" else None)
            elif suffix in {".html", ".htm"}:
                self.extract_html(path, display_name)
            elif suffix in {".eml", ".msg"}:
                self.extract_email(path, display_name)
            elif suffix in {".srt", ".vtt", ".ass", ".ssa"}:
                self.extract_subtitle(path, display_name)
            elif suffix in TEXT_EXTENSIONS:
                self.extract_text(path, display_name)
            else:
                self.extract_unknown(path, display_name)
            self.stats["files_extracted"] += 1
        except JobCancelled:
            raise
        except Exception as exc:  # 单个坏文件不应让整个任务失败。
            message = f"{display_name}：{type(exc).__name__}：{exc}"
            self.warnings.append(message)
            if archive_depth == 0 and display_name not in self.failed_files:
                self.failed_files.append(display_name)
            self.add_section(f"解析提示｜{display_name}", f"该文件未能完整解析。\n原因：{exc}")

    def add_section(self, title: str, body: str) -> None:
        body = body.strip()
        if len(body) > MAX_TEXT_CHARS:
            body = body[:MAX_TEXT_CHARS] + f"\n\n[内容过长，已在 {MAX_TEXT_CHARS:,} 字符处截断]"
        self.sections.append((title, body or "（未提取到可读文本）"))

    def file_metadata(self, path: Path) -> str:
        stat = path.stat()
        mime, encoding = mimetypes.guess_type(path.name)
        return "\n".join(
            [
                f"名称：{path.name}",
                f"扩展名：{path.suffix.lower() or '无'}",
                f"MIME 推测：{mime or '未知'}",
                f"内容编码：{encoding or '无/未知'}",
                f"大小：{human_size(stat.st_size)}（{stat.st_size:,} 字节）",
                f"修改时间：{datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec='seconds')}",
            ]
        )

    def extract_text(self, path: Path, display_name: str) -> None:
        text, encoding = read_text_best_effort(path)
        if path.suffix.lower() in {".json", ".jsonl"}:
            try:
                parsed = json.loads(text)
                text = json.dumps(parsed, ensure_ascii=False, indent=2)
            except json.JSONDecodeError:
                pass
        self.add_section(f"文本内容｜{display_name}", f"检测编码：{encoding}\n\n{text}")

    def extract_html(self, path: Path, display_name: str) -> None:
        raw, encoding = read_text_best_effort(path)
        parser = TextHTMLParser()
        parser.feed(raw)
        title = re.search(r"<title[^>]*>(.*?)</title>", raw, flags=re.I | re.S)
        details = [f"检测编码：{encoding}"]
        if title:
            details.append(f"页面标题：{html.unescape(re.sub('<[^>]+>', '', title.group(1))).strip()}")
        details.extend(["", parser.text()])
        self.add_section(f"网页正文｜{display_name}", "\n".join(details))

    def extract_subtitle(self, path: Path, display_name: str) -> None:
        raw, encoding = read_text_best_effort(path)
        cleaned = re.sub(r"<[^>]+>", "", raw)
        cleaned = re.sub(r"^\s*\d+\s*$", "", cleaned, flags=re.M)
        self.add_section(f"字幕内容｜{display_name}", f"检测编码：{encoding}\n\n{cleaned.strip()}")

    def extract_email(self, path: Path, display_name: str) -> None:
        if path.suffix.lower() == ".msg":
            raise RuntimeError("旧式 Outlook .msg 暂不支持直接解析，建议另存为 .eml 或 PDF")
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        lines = [
            f"主题：{message.get('subject', '')}",
            f"发件人：{message.get('from', '')}",
            f"收件人：{message.get('to', '')}",
            f"抄送：{message.get('cc', '')}",
            f"日期：{message.get('date', '')}",
            "",
        ]
        body = message.get_body(preferencelist=("plain", "html"))
        if body:
            content = body.get_content()
            if body.get_content_type() == "text/html":
                parser = TextHTMLParser()
                parser.feed(content)
                content = parser.text()
            lines.append(content)
        attachments = [part.get_filename() for part in message.iter_attachments() if part.get_filename()]
        if attachments:
            lines.extend(["", "附件：" + "；".join(attachments)])
        self.add_section(f"邮件内容｜{display_name}", "\n".join(lines))

    def extract_media(self, path: Path, display_name: str) -> None:
        if not shutil.which("ffprobe") or not shutil.which("ffmpeg"):
            raise RuntimeError("未找到 FFmpeg/FFprobe，请先安装并加入 PATH")
        data = probe_media(path)
        summary, has_video, has_audio, duration = media_summary(data)
        kind = "视频" if has_video else "音频"
        self.add_section(f"{kind}技术信息｜{display_name}", summary)

        processing_path = path
        processing_duration = duration
        compressed_path: Path | None = None
        keep_compressed_result = False
        try:
            if has_video and self.options.video_compression_enabled:
                compression_dir = self.compressed_keep_dir if self.options.video_keep_compressed else self.compressed_work_dir
                safe_stem = re.sub(r"[^0-9A-Za-z\u4e00-\u9fff._-]+", "_", Path(display_name).stem).strip("._")[:80]
                target_name = f"{safe_stem or 'video'}_{self.next_asset_name('compressed').split('_')[-1]}.mp4"
                compressed_path = compression_dir / target_name
                compressed = compress_video(path, compressed_path, self.options, self.status)
                compressed_data = probe_media(compressed.path)
                compressed_summary, _, compressed_has_audio, compressed_duration = media_summary(compressed_data)
                change_label = (
                    f"减少 {compressed.reduction_percent:.1f}%"
                    if compressed.reduction_percent >= 0
                    else f"增加 {abs(compressed.reduction_percent):.1f}%"
                )
                meets_savings_threshold = (
                    compressed.reduction_percent >= self.options.video_compression_min_savings_percent
                )
                keep_compressed_result = self.options.video_keep_compressed and meets_savings_threshold
                protection_label = (
                    f"达到至少节省 {self.options.video_compression_min_savings_percent}% 的要求"
                    if meets_savings_threshold
                    else f"未达到至少节省 {self.options.video_compression_min_savings_percent}% 的要求，不保留压缩副本"
                )
                self.add_section(
                    f"压缩视频信息｜{display_name}",
                    "\n".join(
                        [
                            f"压缩文件：{compressed.path.name}",
                            f"编码设备：{compressed.encoder}",
                            f"原始大小：{human_size(compressed.original_bytes)}",
                            f"压缩后大小：{human_size(compressed.compressed_bytes)}（{change_label}）",
                            f"解析来源：{'压缩后视频' if self.options.video_parse_source == 'compressed' else '原始视频'}",
                            f"压缩收益保护：{protection_label}",
                            f"压缩文件处理：{'保留并可单独下载' if keep_compressed_result else '解析完成后删除'}",
                            "",
                            compressed_summary,
                        ]
                    ),
                )
                self.stats["compressed_videos"] += 1
                self.stats["compressed_bytes_before"] += compressed.original_bytes
                self.stats["compressed_bytes_after"] += compressed.compressed_bytes
                if keep_compressed_result:
                    self.kept_compressed_files.append(compressed.path)
                elif self.options.video_keep_compressed and not meets_savings_threshold:
                    self.warnings.append(
                        f"{display_name} 的压缩副本仅节省 {compressed.reduction_percent:.1f}%，"
                        f"低于 {self.options.video_compression_min_savings_percent}% 门槛，已自动丢弃。"
                    )
                if self.options.video_parse_source == "compressed":
                    processing_path = compressed.path
                    processing_duration = compressed_duration
                    has_audio = compressed_has_audio

            if has_video and self.options.max_output_images:
                source_label = "压缩后视频" if processing_path == compressed_path else "原始视频"
                self.status(f"正在从{source_label}为 {display_name} 选取代表性画面")
                frame_dir = self.visual_dir / self.next_asset_name("video")
                for frame_path, label in extract_video_frames(processing_path, frame_dir, processing_duration, self.options):
                    friendly_label = label.replace(processing_path.name, display_name, 1)
                    self.visuals.append(VisualAsset(frame_path, friendly_label, display_name, "视频帧"))

            if has_audio and self.options.transcription_enabled:
                source_label = "压缩后视频" if processing_path == compressed_path else kind
                self.status(f"正在准备 {display_name} 的{source_label}音轨")
                audio_path = self.audio_dir / f"{self.next_asset_name('audio')}.wav"
                extract_audio(processing_path, audio_path)
                try:
                    transcript = transcribe_audio(audio_path, self.options, self.status)
                    self.add_section(f"语音转写｜{display_name}", transcript)
                    self.stats["transcripts"] += 1
                except Exception as exc:
                    self.warnings.append(f"{display_name} 的语音转写失败：{exc}")
                    self.add_section(
                        f"语音转写提示｜{display_name}",
                        f"音轨已经成功分离，但语音识别未完成。\n原因：{exc}\n"
                        "如果这是首次运行，请确认网络可用于下载所选 Whisper 模型。",
                    )
        finally:
            if compressed_path and not keep_compressed_result:
                compressed_path.unlink(missing_ok=True)

    def extract_image(self, path: Path, display_name: str) -> None:
        self.status(f"正在分析图片 {display_name}")
        try:
            with Image.open(path) as opened:
                image = ImageOps.exif_transpose(opened)
                metadata = [
                    f"像素尺寸：{image.width} × {image.height}",
                    f"色彩模式：{image.mode}",
                    f"图片格式：{opened.format or path.suffix.lstrip('.')}",
                    f"帧数：{getattr(opened, 'n_frames', 1)}",
                ]
                exif = opened.getexif()
                if exif:
                    selected = []
                    for tag, value in exif.items():
                        if len(selected) >= 20:
                            break
                        if isinstance(value, (str, int, float)) and len(str(value)) < 200:
                            selected.append(f"EXIF {tag}：{value}")
                    metadata.extend(selected)
                target = self.visual_dir / f"{self.next_asset_name('image')}.png"
                image.thumbnail((2400, 2400), Image.Resampling.LANCZOS)
                if image.mode not in {"RGB", "RGBA"}:
                    image = image.convert("RGB")
                image.save(target, "PNG", optimize=True)
            if self.options.max_output_images:
                self.visuals.append(VisualAsset(target, display_name, display_name, "原始图片"))
            self.add_section(f"图片信息｜{display_name}", "\n".join(metadata))
            if self.options.ocr_images:
                ocr_text = self.ocr(target)
                self.add_section(f"图片文字 OCR｜{display_name}", ocr_text or "（未识别到清晰文字）")
        except Exception as exc:
            raise RuntimeError(f"图片解码失败：{exc}") from exc

    def extract_pdf(self, path: Path, display_name: str) -> None:
        from pypdf import PdfReader

        self.status(f"正在提取 PDF {display_name} 的文字和页面")
        reader = PdfReader(str(path))
        metadata = reader.metadata or {}
        lines = [f"页数：{len(reader.pages)}"]
        for key, value in metadata.items():
            if value:
                lines.append(f"{str(key).lstrip('/')}：{value}")
        lines.append("")
        rendered: dict[int, Path] = {}
        selected_pages = select_indices(len(reader.pages), self.options.max_document_pages)
        if self.options.render_documents and self.options.max_output_images:
            rendered = self.render_pdf_pages(path, display_name, selected_pages)

        for index, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text(extraction_mode="layout") or ""
            except TypeError:
                page_text = page.extract_text() or ""
            if not page_text.strip() and index in rendered and self.options.ocr_images:
                page_text = self.ocr(rendered[index])
                if page_text:
                    page_text = "[扫描页 OCR]\n" + page_text
            lines.extend([f"--- 第 {index + 1} 页 ---", page_text.strip() or "（本页未提取到文字）", ""])
        self.add_section(f"PDF 内容｜{display_name}", "\n".join(lines))

    def render_pdf_pages(self, path: Path, display_name: str, page_indices: list[int]) -> dict[int, Path]:
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(path))
        rendered: dict[int, Path] = {}
        for index in page_indices:
            try:
                page = document[index]
                bitmap = page.render(scale=1.5)
                image = bitmap.to_pil()
                target = self.visual_dir / f"{self.next_asset_name('pdf')}.png"
                image.save(target, "PNG", optimize=True)
                rendered[index] = target
                self.visuals.append(VisualAsset(target, f"{display_name} · 第 {index + 1} 页", display_name, "PDF 页面"))
                page.close()
            except Exception as exc:
                self.warnings.append(f"{display_name} 第 {index + 1} 页渲染失败：{exc}")
        document.close()
        return rendered

    def extract_docx(self, path: Path, display_name: str) -> None:
        from docx import Document

        self.status(f"正在读取 Word 文档 {display_name}")
        document = Document(str(path))
        lines: list[str] = []
        props = document.core_properties
        if props.title:
            lines.append(f"标题：{props.title}")
        if props.author:
            lines.append(f"作者：{props.author}")
        lines.append("")
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                style = paragraph.style.name if paragraph.style else ""
                prefix = f"[{style}] " if style and style != "Normal" else ""
                lines.append(prefix + text)
        for table_index, table in enumerate(document.tables, 1):
            lines.extend(["", f"--- 表格 {table_index} ---"])
            for row in table.rows:
                lines.append("\t".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
        for section_index, section in enumerate(document.sections, 1):
            header = " ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
            footer = " ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
            if header:
                lines.append(f"[第 {section_index} 节页眉] {header}")
            if footer:
                lines.append(f"[第 {section_index} 节页脚] {footer}")
        self.add_section(f"Word 内容｜{display_name}", "\n".join(lines))

        if self.options.max_output_images:
            for relation in document.part.rels.values():
                if "image" not in relation.reltype:
                    continue
                part = relation.target_part
                extension = Path(str(part.partname)).suffix or ".bin"
                raw_path = self.visual_dir / f"{self.next_asset_name('docx')}{extension}"
                raw_path.write_bytes(part.blob)
                self.add_embedded_image(raw_path, display_name, "Word 内嵌图片")

    def extract_pptx(self, path: Path, display_name: str) -> None:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        self.status(f"正在读取演示文稿 {display_name}")
        presentation = Presentation(str(path))
        lines = [f"幻灯片数量：{len(presentation.slides)}", ""]
        for slide_index, slide in enumerate(presentation.slides, 1):
            lines.append(f"--- 第 {slide_index} 页 ---")
            slide_text: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    value = shape.text.strip()
                    if value:
                        slide_text.append(value)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        slide_text.append("\t".join(cell.text.replace("\n", " / ").strip() for cell in row.cells))
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and self.options.max_output_images:
                    image = shape.image
                    raw_path = self.visual_dir / f"{self.next_asset_name('pptx')}.{image.ext}"
                    raw_path.write_bytes(image.blob)
                    self.add_embedded_image(raw_path, display_name, f"PPT 第 {slide_index} 页内嵌图片")
                if getattr(shape, "has_chart", False):
                    slide_text.append(f"[图表：{getattr(shape, 'name', '未命名')}] ")
            try:
                notes_frame = slide.notes_slide.notes_text_frame
                notes = notes_frame.text.strip() if notes_frame else ""
                if notes:
                    slide_text.append("[演讲者备注]\n" + notes)
            except (AttributeError, ValueError):
                pass
            lines.append("\n".join(slide_text) or "（本页未发现可提取文字）")
            lines.append("")
        self.add_section(f"演示文稿内容｜{display_name}", "\n".join(lines))

    def extract_xlsx(self, path: Path, display_name: str) -> None:
        from openpyxl import load_workbook

        self.status(f"正在读取工作簿 {display_name}")
        workbook = load_workbook(str(path), read_only=True, data_only=False, keep_links=False)
        lines = [f"工作表数量：{len(workbook.sheetnames)}", f"工作表：{'；'.join(workbook.sheetnames)}", ""]
        for worksheet in workbook.worksheets:
            lines.append(f"--- 工作表：{worksheet.title}（{worksheet.max_row} 行 × {worksheet.max_column} 列）---")
            rows_written = 0
            for row in worksheet.iter_rows(
                min_row=1,
                max_row=min(worksheet.max_row or 1, self.options.table_max_rows),
                max_col=min(worksheet.max_column or 1, self.options.table_max_columns),
                values_only=True,
            ):
                lines.append("\t".join(format_cell(value) for value in row))
                rows_written += 1
            if (worksheet.max_row or 0) > rows_written:
                lines.append(f"[该表其余 {worksheet.max_row - rows_written} 行已省略，可提高‘每表最大行数’后重试]")
            if (worksheet.max_column or 0) > self.options.table_max_columns:
                lines.append(f"[右侧列已在第 {self.options.table_max_columns} 列处截断]")
            lines.append("")
        workbook.close()
        self.add_section(f"工作簿内容｜{display_name}", "\n".join(lines))

    def extract_delimited(self, path: Path, display_name: str, delimiter: str | None) -> None:
        text, encoding = read_text_best_effort(path)
        sample = text[:8192]
        if delimiter is None:
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                delimiter = ","
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        lines = [f"检测编码：{encoding}", f"分隔符：{repr(delimiter)}", ""]
        for row_index, row in enumerate(reader):
            if row_index >= self.options.table_max_rows:
                lines.append(f"[已在第 {self.options.table_max_rows} 行处截断]")
                break
            lines.append("\t".join(row[: self.options.table_max_columns]))
        self.add_section(f"表格内容｜{display_name}", "\n".join(lines))

    def extract_archive(self, path: Path, display_name: str, depth: int) -> None:
        self.stats["archives"] += 1
        if depth >= self.options.archive_depth:
            self.add_section(
                f"压缩包提示｜{display_name}",
                f"已达到压缩包递归深度上限 {self.options.archive_depth}，未继续展开。",
            )
            return
        self.status(f"正在安全展开压缩包 {display_name}")
        target = self.archive_dir / self.next_asset_name("archive")
        target.mkdir(parents=True, exist_ok=True)
        names: list[str] = []
        lower_name = path.name.lower()
        if path.suffix.lower() == ".zip":
            names = self.safe_extract_zip(path, target)
        elif path.suffix.lower() == ".7z":
            names = self.safe_extract_7z(path, target)
        elif lower_name.endswith((".tar", ".tar.gz", ".tgz", ".tar.bz2", ".tar.xz")):
            names = self.safe_extract_tar(path, target)
        elif path.suffix.lower() in {".gz", ".bz2", ".xz"} and not lower_name.endswith((".tar.gz", ".tar.bz2", ".tar.xz")):
            output_name = Path(path.stem).name or "decompressed"
            output_path = target / output_name
            opener = {".gz": gzip.open, ".bz2": bz2.open, ".xz": lzma.open}[path.suffix.lower()]
            with opener(path, "rb") as source, output_path.open("wb") as destination:
                copy_with_limit(source, destination, MAX_ARCHIVE_BYTES)
            names = [output_name]
        else:
            raise RuntimeError("该压缩格式当前无法安全展开；建议转换为 ZIP、7Z 或 TAR")

        file_paths = [item for item in target.rglob("*") if item.is_file()]
        file_paths = file_paths[: self.options.max_archive_files]
        inventory = [f"成员总数（目录与文件）：{len(names)}", f"本次解析文件数：{len(file_paths)}", ""]
        inventory.extend(f"- {item}" for item in names[:500])
        if len(names) > 500:
            inventory.append(f"- …另有 {len(names) - 500} 项未列出")
        if len(file_paths) >= self.options.max_archive_files:
            inventory.append(f"\n[已达到单个压缩包 {self.options.max_archive_files} 个文件的解析上限]")
        self.add_section(f"压缩包目录｜{display_name}", "\n".join(inventory))

        for member in file_paths:
            relative = member.relative_to(target).as_posix()
            self.process_file(member, f"{display_name} / {relative}", depth + 1)

    def safe_extract_zip(self, path: Path, target: Path) -> list[str]:
        names: list[str] = []
        total = 0
        with zipfile.ZipFile(path) as archive:
            for info in archive.infolist():
                if len(names) >= self.options.max_archive_files * 5:
                    break
                names.append(info.filename)
                if info.is_dir():
                    continue
                total += info.file_size
                if total > MAX_ARCHIVE_BYTES:
                    raise RuntimeError("压缩包展开后超过 4 GB 安全上限")
                destination = safe_member_path(target, info.filename)
                destination.parent.mkdir(parents=True, exist_ok=True)
                with archive.open(info) as source, destination.open("wb") as output:
                    shutil.copyfileobj(source, output, length=1024 * 1024)
        return names

    def safe_extract_tar(self, path: Path, target: Path) -> list[str]:
        names: list[str] = []
        total = 0
        with tarfile.open(path, "r:*") as archive:
            for member in archive:
                if len(names) >= self.options.max_archive_files * 5:
                    break
                names.append(member.name)
                if not member.isfile():
                    continue
                total += member.size
                if total > MAX_ARCHIVE_BYTES:
                    raise RuntimeError("压缩包展开后超过 4 GB 安全上限")
                source = archive.extractfile(member)
                if source is None:
                    continue
                destination = safe_member_path(target, member.name)
                destination.parent.mkdir(parents=True, exist_ok=True)
                with source, destination.open("wb") as output:
                    shutil.copyfileobj(source, output, length=1024 * 1024)
        return names

    def safe_extract_7z(self, path: Path, target: Path) -> list[str]:
        try:
            import py7zr
        except ImportError as exc:
            raise RuntimeError("未安装 py7zr，无法展开 7Z 文件") from exc
        with py7zr.SevenZipFile(path, mode="r") as archive:
            file_info = archive.list()
            names = [item.filename for item in file_info]
            if len(names) > self.options.max_archive_files * 5:
                raise RuntimeError("7Z 成员过多，已超过安全上限")
            if sum(int(item.uncompressed or 0) for item in file_info) > MAX_ARCHIVE_BYTES:
                raise RuntimeError("7Z 展开后超过 4 GB 安全上限")
            for name in names:
                safe_member_path(target, name)
            archive.extractall(path=target)
        return names

    def extract_unknown(self, path: Path, display_name: str) -> None:
        try:
            data = probe_media(path)
            if data.get("streams"):
                self.extract_media(path, display_name)
                return
        except Exception:
            pass
        raw = path.read_bytes()[:65536]
        if raw:
            try:
                text, encoding = decode_bytes(raw)
                printable_ratio = sum(char.isprintable() or char in "\r\n\t" for char in text) / max(1, len(text))
                if printable_ratio > 0.85:
                    self.add_section(f"文本内容（自动识别）｜{display_name}", f"检测编码：{encoding}\n\n{text}")
                    return
            except Exception:
                pass
        self.add_section(
            f"未知格式｜{display_name}",
            "已保留文件层面的元数据，但没有可靠的内容解析器。建议将该文件转换为 PDF、文本、图片、Office Open XML 或 FFmpeg 可识别的媒体格式后重试。",
        )

    def add_embedded_image(self, raw_path: Path, source_name: str, kind: str) -> None:
        try:
            with Image.open(raw_path) as opened:
                image = ImageOps.exif_transpose(opened)
                image.thumbnail((2400, 2400), Image.Resampling.LANCZOS)
                if image.mode not in {"RGB", "RGBA"}:
                    image = image.convert("RGB")
                target = self.visual_dir / f"{self.next_asset_name('embedded')}.png"
                image.save(target, "PNG", optimize=True)
            self.visuals.append(VisualAsset(target, kind, source_name, kind))
        except Exception as exc:
            self.warnings.append(f"{source_name} 中有一张内嵌图片无法解码：{exc}")

    def ocr(self, path: Path) -> str:
        try:
            if self._ocr_engine is None:
                from rapidocr_onnxruntime import RapidOCR

                self._ocr_engine = RapidOCR()
            result, _ = self._ocr_engine(str(path))
            if not result:
                return ""
            lines = []
            for item in result:
                if len(item) >= 3 and float(item[2]) >= 0.45:
                    lines.append(str(item[1]).strip())
            return "\n".join(line for line in lines if line)
        except Exception as exc:
            self.warnings.append(f"OCR 未能处理 {path.name}：{exc}")
            return ""

    def next_asset_name(self, prefix: str) -> str:
        self._asset_counter += 1
        return f"{prefix}_{self._asset_counter:05d}"


def safe_member_path(root: Path, member_name: str) -> Path:
    normalized = member_name.replace("\\", "/")
    if normalized.startswith("/") or re.match(r"^[A-Za-z]:", normalized):
        raise RuntimeError(f"压缩包包含不安全的绝对路径：{member_name}")
    destination = (root / normalized).resolve()
    root_resolved = root.resolve()
    if not destination.is_relative_to(root_resolved):
        raise RuntimeError(f"压缩包包含越界路径：{member_name}")
    return destination


def copy_with_limit(source: object, destination: object, byte_limit: int) -> int:
    total = 0
    while True:
        chunk = source.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > byte_limit:
            raise RuntimeError(f"解压内容超过 {byte_limit / 1024**3:.0f} GB 安全上限")
        destination.write(chunk)
    return total


def read_text_best_effort(path: Path) -> tuple[str, str]:
    return decode_bytes(path.read_bytes()[: MAX_TEXT_CHARS * 4])


def decode_bytes(raw: bytes) -> tuple[str, str]:
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16"), "UTF-16"
    for encoding in ("utf-8-sig", "gb18030", "big5", "shift_jis"):
        try:
            return raw.decode(encoding), encoding.upper()
        except UnicodeDecodeError:
            continue
    try:
        from charset_normalizer import from_bytes

        match = from_bytes(raw).best()
        if match:
            return str(match), match.encoding or "自动检测"
    except Exception:
        pass
    return raw.decode("utf-8", errors="replace"), "UTF-8（含替换字符）"


def select_indices(total: int, maximum: int) -> list[int]:
    if total <= 0:
        return []
    if total <= maximum:
        return list(range(total))
    if maximum == 1:
        return [0]
    return sorted({round(index * (total - 1) / (maximum - 1)) for index in range(maximum)})


def format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    return str(value).replace("\r", " ").replace("\n", " / ").replace("\t", " ")
