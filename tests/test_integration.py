from __future__ import annotations

import io
import json
import hashlib
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches

from app.main import app


class EndToEndTests(unittest.TestCase):
    def test_resumable_upload_verifies_chunks_and_waits_for_preflight_confirmation(self) -> None:
        content = (b"contextkit-resumable-upload\n" * 700_000) + b"done"
        chunk_size = 16 * 1024 * 1024
        options = {"transcription_enabled": False, "ocr_images": False, "max_output_images": 0}
        payload = {
            "files": [
                {
                    "name": "notes.txt",
                    "relative_path": "study/notes.txt",
                    "size": len(content),
                    "content_type": "text/plain",
                    "last_modified": 123456,
                }
            ],
            "options": options,
        }

        with TestClient(app) as client:
            created = client.post("/api/uploads", json=payload)
            self.assertEqual(created.status_code, 201, created.text)
            job = created.json()
            job_id = job["id"]
            remote = job["upload"]["files"][0]
            self.assertEqual(job["state"], "uploading")
            self.assertEqual(remote["total_chunks"], 2)

            second = content[chunk_size:]
            bad_hash = client.put(
                f"/api/uploads/{job_id}/files/{remote['id']}/chunks/1",
                content=second,
                headers={"X-Chunk-SHA256": "0" * 64},
            )
            self.assertEqual(bad_hash.status_code, 409)

            for index, chunk in ((1, second), (0, content[:chunk_size])):
                response = client.put(
                    f"/api/uploads/{job_id}/files/{remote['id']}/chunks/{index}",
                    content=chunk,
                    headers={"X-Chunk-SHA256": hashlib.sha256(chunk).hexdigest()},
                )
                self.assertEqual(response.status_code, 200, response.text)

            status = client.get(f"/api/uploads/{job_id}").json()["upload"]
            self.assertEqual(status["progress"], 100)
            self.assertEqual(status["files"][0]["uploaded_chunks"], [0, 1])

            ready = client.post(f"/api/uploads/{job_id}/complete")
            self.assertEqual(ready.status_code, 200, ready.text)
            ready_job = ready.json()
            self.assertEqual(ready_job["state"], "ready")
            self.assertEqual(ready_job["preflight"]["summary"]["file_count"], 1)
            self.assertEqual(ready_job["preflight"]["files"][0]["name"], "study/notes.txt")

            started = client.post(f"/api/jobs/{job_id}/start")
            self.assertEqual(started.status_code, 202, started.text)
            completed = client.get(f"/api/jobs/{job_id}").json()
            self.assertEqual(completed["state"], "completed")
            self.assertEqual(client.get(f"/api/jobs/{job_id}/download").status_code, 200)

            retry = client.post(f"/api/jobs/{job_id}/retry")
            self.assertEqual(retry.status_code, 202, retry.text)
            retry_id = retry.json()["id"]
            self.assertEqual(client.get(f"/api/jobs/{retry_id}").json()["source_job_id"], job_id)
            self.assertEqual(client.delete(f"/api/jobs/{retry_id}").status_code, 204)
            self.assertEqual(client.delete(f"/api/jobs/{job_id}").status_code, 204)

    def test_upload_can_be_cancelled_then_cleaned(self) -> None:
        payload = {
            "files": [{"name": "pause.txt", "size": 100, "content_type": "text/plain", "last_modified": 1}],
            "options": {"transcription_enabled": False},
        }
        with TestClient(app) as client:
            created = client.post("/api/uploads", json=payload)
            job_id = created.json()["id"]
            cancelled = client.post(f"/api/jobs/{job_id}/cancel")
            self.assertEqual(cancelled.status_code, 200)
            self.assertEqual(cancelled.json()["state"], "cancelled")
            self.assertEqual(client.delete(f"/api/jobs/{job_id}").status_code, 204)

    @unittest.skipUnless(shutil.which("ffmpeg") and shutil.which("ffprobe"), "FFmpeg is required")
    def test_mixed_upload_produces_strict_context_pack(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            fixture_dir = Path(temp)
            payloads = self.build_fixtures(fixture_dir)
            options = {
                "transcription_enabled": False,
                "ocr_images": False,
                "video_frame_count": 4,
                "video_compression_enabled": True,
                "video_compression_device": "cpu",
                "video_compression_max_height": 360,
                "video_compression_quality": 28,
                "video_parse_source": "compressed",
                "video_keep_compressed": True,
                "max_output_images": 3,
                "frames_per_sheet": 4,
                "max_document_pages": 2,
            }
            files = [("files", (name, content, mime)) for name, content, mime in payloads]

            with TestClient(app) as client:
                response = client.post("/api/jobs", files=files, data={"options": json.dumps(options)})
                self.assertEqual(response.status_code, 202, response.text)
                job_id = response.json()["id"]

                status = client.get(f"/api/jobs/{job_id}")
                self.assertEqual(status.status_code, 200)
                self.assertEqual(status.json()["state"], "completed", status.text)
                compressed_files = status.json()["compressed_files"]
                self.assertEqual(len(compressed_files), 1)
                self.assertTrue(compressed_files[0]["name"].endswith(".mp4"))

                compressed = client.get(f"/api/jobs/{job_id}/compressed/{compressed_files[0]['name']}")
                self.assertEqual(compressed.status_code, 200)
                self.assertGreater(len(compressed.content), 0)

                download = client.get(f"/api/jobs/{job_id}/download")
                self.assertEqual(download.status_code, 200)
                with zipfile.ZipFile(io.BytesIO(download.content)) as archive:
                    names = archive.namelist()
                    self.assertIn("context.txt", names)
                    png_names = [name for name in names if name.endswith(".png")]
                    self.assertLessEqual(len(png_names), 3)
                    self.assertTrue(all(name == "context.txt" or name.endswith(".png") for name in names))
                    report = archive.read("context.txt").decode("utf-8-sig")
                    for expected in ("notes.txt", "photo.png", "report.docx", "slides.pptx", "table.xlsx", "bundle.zip", "clip.mp4"):
                        self.assertIn(expected, report)
                    self.assertIn("视频技术信息", report)
                    self.assertIn("压缩包目录", report)
                    self.assertIn("演示文稿内容", report)
                    self.assertIn("压缩视频信息", report)
                    self.assertIn("压缩后视频", report)

                deleted = client.delete(f"/api/jobs/{job_id}")
                self.assertEqual(deleted.status_code, 204)

    @staticmethod
    def build_fixtures(root: Path) -> list[tuple[str, bytes, str]]:
        image_path = root / "photo.png"
        image = Image.new("RGB", (640, 360), "#e9e4d8")
        drawing = ImageDraw.Draw(image)
        drawing.rectangle((70, 70, 570, 290), outline="#e66346", width=8)
        drawing.text((210, 165), "ContextKit", fill="#17201d")
        image.save(image_path)

        pdf_path = root / "scan.pdf"
        image.save(pdf_path, "PDF", resolution=100)

        docx_path = root / "report.docx"
        document = Document()
        document.add_heading("项目周报", level=1)
        document.add_paragraph("本周完成了多模态文件解析与输出验证。")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "项目"
        table.cell(0, 1).text = "状态"
        table.cell(1, 0).text = "视频抽帧"
        table.cell(1, 1).text = "完成"
        document.save(docx_path)

        pptx_path = root / "slides.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "多模态信息抽取"
        slide.placeholders[1].text = "语音转写\n视觉拼图\n文档结构"
        slide.shapes.add_picture(str(image_path), Inches(5.2), Inches(2), width=Inches(3.8))
        presentation.save(pptx_path)

        xlsx_path = root / "table.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "里程碑"
        sheet.append(["任务", "负责人", "完成度"])
        sheet.append(["解析器", "Local", 1])
        sheet.append(["界面", "Local", 0.95])
        workbook.save(xlsx_path)

        archive_buffer = io.BytesIO()
        with zipfile.ZipFile(archive_buffer, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("inside/readme.md", "# 压缩包内部文档\n这是递归解析测试。")
            archive.writestr("inside/data.csv", "name,value\nalpha,1\nbeta,2\n")

        video_path = root / "clip.mp4"
        command = [
            "ffmpeg", "-hide_banner", "-loglevel", "error", "-y",
            "-f", "lavfi", "-i", "testsrc2=size=640x360:rate=12:duration=2",
            "-f", "lavfi", "-i", "sine=frequency=440:sample_rate=16000:duration=2",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", str(video_path),
        ]
        subprocess.run(command, check=True, capture_output=True)

        return [
            ("notes.txt", "这是一个端到端测试文件。".encode("utf-8"), "text/plain"),
            ("photo.png", image_path.read_bytes(), "image/png"),
            ("scan.pdf", pdf_path.read_bytes(), "application/pdf"),
            ("report.docx", docx_path.read_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ("slides.pptx", pptx_path.read_bytes(), "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
            ("table.xlsx", xlsx_path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("bundle.zip", archive_buffer.getvalue(), "application/zip"),
            ("clip.mp4", video_path.read_bytes(), "video/mp4"),
        ]


if __name__ == "__main__":
    unittest.main()
